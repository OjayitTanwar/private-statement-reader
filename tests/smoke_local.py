"""Real synthetic scan -> OCR -> rows -> HTTP review/export test. No private inputs."""
import hashlib,io,json,os,sys,threading,time,urllib.request,urllib.error,uuid
from pathlib import Path
from PIL import Image,ImageDraw,ImageFont,ImageEnhance,ImageFilter

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
os.environ['STATEMENT_DATA']=str(ROOT/'work/smoke-data')
os.environ['STATEMENT_VISION']=str(ROOT/'outputs/Private Statement Reader.app/Contents/Resources/VisionOCR')
import app


def make_pdf(path):
    font=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf',30)
    title=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial Bold.ttf',42)
    images=[]
    expected=[('01/04/2025','UPI/AJAI SHARMA/123456','1,000.00','','9,000.00'),
              ('02/04/2025','NEFT/RECEIVED/RAMESH','','5,000.00','14,000.00'),
              ('03/04/2025','ATM CASH','2,500.25','','11,499.75'),
              ('04/04/2025','IMPS/MOHAMMAD KHAN','','200.00','11,699.75'),
              ('05/04/2025','POS LOCAL STORE','199.75','','11,500.00'),
              ('06/04/2025','NEFT/RECEIVED/SHARMA','','1,23,456.70','1,34,956.70')]
    for page in range(2):
        im=Image.new('RGB',(2200,2800),'white');draw=ImageDraw.Draw(im)
        draw.text((80,80),'SYNTHETIC BANK — Account Statement',font=title,fill='#202020')
        draw.text((80,160),'Account Number: 0000000000   IFSC: TEST0000001',font=font,fill='#202020')
        draw.text((80,220),'TEST DATA ONLY — not a real bank account',font=font,fill='#303030')
        for text,x in [('Date',80),('Particulars',360),('Withdrawal',1320),('Deposit',1590),('Balance',1870)]:draw.text((x,350),text,font=font,fill='black')
        if page==0:
            draw.text((360,410),'Opening Balance',font=font,fill='black');draw.text((1870,410),'10,000.00',font=font,fill='black')
        for n,record in enumerate(expected[page*3:page*3+3]):
            for text,x in zip(record,(80,360,1320,1590,1870)):
                if text:draw.text((x,500+n*120),text,font=font,fill='#252525')
        draw.text((80,2600),f'Synthetic statement page {page+1} of 2',font=font,fill='#404040')
        if page==1:
            im=ImageEnhance.Contrast(im).enhance(.45).filter(ImageFilter.GaussianBlur(.45))
        images.append(im)
    images[0].save(path,'PDF',save_all=True,append_images=images[1:],resolution=220)
    return expected


def main():
    app.init_db();doc=str(uuid.uuid4());folder=app.DOCS/doc;folder.mkdir();expected=make_pdf(folder/'source.pdf')
    source_hash=hashlib.sha256((folder/'source.pdf').read_bytes()).hexdigest()
    with app.connection() as c:
        c.execute('INSERT INTO bank_documents(id,name,status,created,config) VALUES(?,?,?,?,?)',(doc,'Synthetic Indian bank statement.pdf','queued',app.stamp(),json.dumps({'thorough':True,'opening':None,'order':'auto'})))
    started=time.monotonic();app.process(doc);d=app.document(doc);rows=app.rows_for(doc)
    report={'status':d['status'],'elapsed_seconds':round(time.monotonic()-started,1),'rows':len(rows),'report':d['report'],'config':d['config'],'extracted':[{k:r.get(k) for k in ('date','debit','credit','balance','confidence','reason','_engine')} for r in rows]}
    (ROOT/'work/smoke-result.json').write_text(json.dumps(report,indent=2))
    print(json.dumps(report,indent=2),flush=True)
    assert d['status']=='ready',d
    assert len(rows)==6,len(rows)
    for row,exp in zip(rows,expected):
        assert row['balance']==app.parse_amount(exp[4]),row
        assert row['debit']==app.parse_amount(exp[2]),row
        assert row['credit']==app.parse_amount(exp[3]),row
        assert (folder/row['crop']).is_file()
        assert 'balance mismatch' not in row['reason'],row
    assert source_hash==hashlib.sha256((folder/'source.pdf').read_bytes()).hexdigest()
    assert d['report']['searchable']
    for pg in (1,2):
        for name in ('tesseract','tesseract-enhanced','apple'):
            assert len(json.loads((folder/f'ocr-{pg}-{name}.json').read_text()))>20, f'{name} did not contribute real OCR'
    extracted=app.run(['pdftotext',folder/'searchable.pdf','-']).stdout
    assert 'Statement' in extracted
    server=app.ThreadingHTTPServer(('127.0.0.1',0),app.Handler)
    threading.Thread(target=server.serve_forever,daemon=True).start();base=f'http://127.0.0.1:{server.server_port}'
    def request(path,payload=None,auth=True):
        headers={'X-Statement-Token':app.TOKEN} if auth else {}
        data=None if payload is None else json.dumps(payload).encode()
        return urllib.request.urlopen(urllib.request.Request(base+path,data=data,headers=headers))
    try:
        try:request('/api/documents',auth=False);raise AssertionError('Unauthenticated API allowed')
        except urllib.error.HTTPError as e:assert e.code==403
        resources=json.load(request('/api/resources'))
        assert resources['memory_limit_kind']=='best_effort'
        changed=json.load(request('/api/resources',{'preset':'light'}))
        assert changed['settings']['documents']==1
        assert json.load(request('/api/resources'))['settings']==changed['settings']
        try:request('/api/resources',{'preset':'custom','documents':0});raise AssertionError('Invalid resource settings allowed')
        except urllib.error.HTTPError as e:assert e.code==400
        assert json.load(request('/api/resources'))['settings']==changed['settings']
        hostile=urllib.request.Request(base+'/api/resources',data=b'{"preset":"maximum"}',headers={'X-Statement-Token':app.TOKEN,'Origin':'https://example.invalid'})
        try:urllib.request.urlopen(hostile);raise AssertionError('Cross-origin resource update allowed')
        except urllib.error.HTTPError as e:assert e.code==403
        assert len(json.load(request('/api/rows/'+doc)))==6
        request('/api/people',[{'name':'Ajay Sharma','aliases':['Ajai Sharma']}])
        assert json.load(request('/api/rows/'+doc))[0]['people_suggestions'][0]['name']=='Ajay Sharma'
        request('/api/row/'+rows[0]['id'],{'debit':1001,'normalized_description':'=BADFORMULA()'})
        corrected=json.load(request('/api/rows/'+doc));assert 'balance mismatch' in corrected[0]['reason']
        assert corrected[0]['original_description']==rows[0]['original_description']
        csv=request('/export/'+doc+'/csv').read().decode('utf-8-sig');assert "'=BADFORMULA()" in csv
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(request('/export/'+doc+'/xlsx').read()));assert wb.active.max_row==7
        request('/api/row/'+rows[0]['id'],{'debit':1000,'state':'approved'})
        audit=json.load(request('/export/'+doc+'/audit'));assert len(audit)>=2
        request('/api/row/'+rows[0]['id'],{'state':'deleted'})
        assert len(request('/export/'+doc+'/csv').read().decode('utf-8-sig').splitlines())==6
        print('PASS: real OCR, searchable PDF, source hashes, crops, names, edits, reconciliation, audit, CSV/XLSX, persistent resource settings and authentication',flush=True)
    finally:server.shutdown()


if __name__=='__main__':main()
