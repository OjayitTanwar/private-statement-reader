import copy
import csv
import io
import json
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from automation import (atomic_json, empty_summary, export_bytes, prepared_rows,
                        write_document_outputs)


def transaction(identifier='row1', **changes):
    row = dict(id=identifier, date='2025-01-01', original_description='=SUM(A1:A2)',
               normalized_description='TRANSFER', counterparty='+Potential formula',
               cheque_number='000012', debit=0.01, credit=0, balance=999.99,
               page=1, confidence=.8, suspicious=True, reason='Uncertain source',
               state='pending', _ocr_before_ai={'balance':999.98},
               _ai_proposal={'balance':999.99},
               _inferred_fields={'balance':{'original_value':999.98,
                                          'not_printed_evidence':True}})
    row.update(changes)
    return row


class AutomaticOutputTests(unittest.TestCase):
    def test_exports_preserve_cents_provenance_and_literal_text(self):
        source=[transaction(), transaction('deleted', state='deleted')]
        before=copy.deepcopy(source)
        rows=prepared_rows(source)
        self.assertEqual(len(rows),1)
        self.assertEqual(source,before)
        self.assertEqual(rows[0]['inference_status'],'Contains inferred values')
        self.assertEqual(json.loads(rows[0]['original_ocr_values']),{'balance':999.98})
        self.assertEqual(json.loads(rows[0]['ai_proposal']),{'balance':999.99})
        self.assertTrue(json.loads(rows[0]['inferred_fields'])['balance']['not_printed_evidence'])
        metadata={'name':'=Untrusted document','report':{'warnings':['=Untrusted note']}}
        data=export_bytes('csv',rows,metadata)
        exported=list(csv.DictReader(io.StringIO(data.decode('utf-8-sig'))))
        self.assertEqual(len(exported),1)
        self.assertEqual(Decimal(exported[0]['debit']),Decimal('0.01'))
        self.assertEqual(Decimal(exported[0]['balance']),Decimal('999.99'))
        self.assertEqual(exported[0]['cheque_number'],'000012')
        self.assertEqual(exported[0]['original_description'],"'=SUM(A1:A2)")
        self.assertEqual(exported[0]['counterparty'],"'+Potential formula")
        from openpyxl import load_workbook
        workbook=load_workbook(io.BytesIO(export_bytes('xlsx',rows,metadata)))
        sheet=workbook['Transactions']
        headers=[cell.value for cell in sheet[1]]
        values={key:sheet.cell(2,index+1) for index,key in enumerate(headers)}
        self.assertEqual(Decimal(str(values['balance'].value)),Decimal('999.99'))
        self.assertEqual(values['balance'].number_format,'#,##0.00')
        self.assertEqual(values['original_description'].data_type,'s')
        self.assertEqual(values['cheque_number'].value,'000012')
        self.assertEqual(workbook['Extraction notes']['B1'].data_type,'s')
        self.assertEqual(workbook['Extraction notes']['B4'].data_type,'s')

    def test_summary_detection_requires_explicit_empty_statement_summary(self):
        self.assertTrue(empty_summary('SUMMARY OF ACCOUNTS\nNo transactions found for this period'))
        self.assertTrue(empty_summary('Summary of Account   No transaction in this period'))
        self.assertFalse(empty_summary('Summary of Accounts\n10 transactions found'))
        self.assertFalse(empty_summary('No transactions found in your search'))
        self.assertFalse(empty_summary('Transfer narration: NO TRANSACTIONS FOR invoice'))

    def test_document_output_keeps_source_and_has_no_partial_artifacts(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);source=root/'documents'/'document1'
            source.mkdir(parents=True)
            original=b'Original private source bytes';searchable=b'Searchable output bytes'
            (source/'original.pdf').write_bytes(original)
            (source/'searchable.pdf').write_bytes(searchable)
            document={'id':'document1','name':'Bank.pdf','status':'ready',
                      'report':{'searchable':True,'warnings':[]}}
            folder=write_document_outputs(root,document,[transaction()])
            self.assertEqual((source/'original.pdf').read_bytes(),original)
            self.assertEqual((folder/'searchable-statement.pdf').read_bytes(),searchable)
            self.assertTrue((folder/'transactions.csv').exists())
            self.assertTrue((folder/'transactions.xlsx').exists())
            self.assertEqual(json.loads((folder/'extraction-report.json').read_text()),document['report'])
            self.assertFalse(list(root.rglob('*.pending')))
            checkpoint=root/'checkpoint.json'
            atomic_json(checkpoint,{'next_page':2,'rows':[transaction()]})
            atomic_json(checkpoint,{'next_page':3,'rows':[]})
            self.assertEqual(json.loads(checkpoint.read_text())['next_page'],3)
            self.assertFalse(list(root.rglob('*.pending')))

    def test_unrecognized_documents_write_report_without_transaction_exports(self):
        with tempfile.TemporaryDirectory() as directory:
            document={'id':'unrecognized','name':'Other.pdf','status':'needs_layout',
                      'report':{'warnings':['Bank table not detected']}}
            folder=write_document_outputs(directory,document,[transaction()])
            self.assertTrue((folder/'extraction-report.json').exists())
            self.assertFalse((folder/'transactions.csv').exists())


if __name__=='__main__':unittest.main()
