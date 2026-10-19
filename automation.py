"""Durable automatic processing and local deliverables, independent of the UI."""
from pathlib import Path
import csv,io,json,os,re
from review_checks import classify_review

FIELDS=['date','original_description','normalized_description','counterparty','cheque_number','debit','credit','balance','page','confidence','suspicious','reason','state','inference_status','original_ocr_values','ai_proposal','inferred_fields']
FIELDS+=['review_category','review_required','review_reason','information_notes','legacy_suspicious']


def atomic_json(path,value):
    path=Path(path);temporary=path.with_suffix(path.suffix+'.pending')
    with temporary.open('w') as f:
        json.dump(value,f,ensure_ascii=False);f.flush();os.fsync(f.fileno())
    temporary.replace(path)


def empty_summary(text):
    text=re.sub(r'\s+',' ',text)
    return bool(re.search(r'\bno\s+transactions?\s+(?:found|for|in)\b',text,re.I)
                and re.search(r'\bsummary\s+of\s+accounts?\b',text,re.I))


def prepared_rows(rows):
    result=[]
    for original in rows:
        if original.get('state')=='deleted':continue
        r=dict(original)
        review=classify_review(original)
        r.update(review_category=review['category'],review_required=review['review_required'],
                 review_reason='; '.join(review['financial_reasons']+['Estimated '+key for key in review['estimated_fields']]),
                 information_notes='; '.join(review['information']),legacy_suspicious=bool(original.get('suspicious')),
                 suspicious=review['review_required'])
        r['inference_status']='AI assisted' if r.get('_ai_proposal') else 'OCR only'
        if r.get('_inferred_fields'):r['inference_status']='Contains inferred values'
        for field,key in [('original_ocr_values','_ocr_before_ai'),('ai_proposal','_ai_proposal'),('inferred_fields','_inferred_fields')]:
            r[field]=json.dumps(r.get(key,{}),ensure_ascii=False)
        result.append(r)
    return result


def safe_cell(value):
    return "'"+value if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@','\t','\r')) else value


def export_bytes(kind,rows,metadata,fields=None):
    fields=fields or FIELDS
    if kind=='csv':
        out=io.StringIO();writer=csv.writer(out);writer.writerow(fields)
        for r in rows:writer.writerow([safe_cell(r.get(k,'')) for k in fields])
        return ('\ufeff'+out.getvalue()).encode('utf-8')
    if kind!='xlsx':raise ValueError('Unsupported transaction format')
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    wb=Workbook();ws=wb.active;ws.title='Transactions';ws.append(fields)
    for r in rows:
        ws.append([safe_cell(r.get(k,'')) for k in fields])
        if r.get('suspicious'):
            for cell in ws[ws.max_row]:cell.fill=PatternFill('solid',fgColor='FFF1C7')
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='174D3C')
    for i,field in enumerate(fields,1):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(i)].width=55 if 'description' in field else 30 if field in ('reason','source_document') else 18
        if field in ('debit','credit','balance'):
            for cells in ws.iter_rows(min_row=2,min_col=i,max_col=i):cells[0].number_format='#,##0.00'
    notes=wb.create_sheet('Extraction notes');notes.append(['Source',safe_cell(metadata.get('name','Batch'))])
    notes.append(['Confidence','OCR quality score; arithmetic agreement does not prove transcription accuracy.'])
    notes.append(['Estimates','Inferred values, AI proposals and original readings are included in the transaction columns.'])
    for warning in metadata.get('report',{}).get('warnings',[]):notes.append(['Note',safe_cell(str(warning))])
    buf=io.BytesIO();wb.save(buf);return buf.getvalue()


def atomic_bytes(path,value):
    path=Path(path);temporary=path.with_suffix(path.suffix+'.pending')
    temporary.write_bytes(value);temporary.replace(path)


def write_document_outputs(data,document,rows):
    stem=re.sub(r'[^\w .-]','_',Path(document['name']).stem)[:100].strip(' .') or 'statement'
    folder=Path(data)/'exports'/f"{stem}--{document['id'][:8]}";folder.mkdir(parents=True,exist_ok=True)
    atomic_json(folder/'extraction-report.json',document['report'])
    if document['status']=='ready':
        prepared=prepared_rows(rows)
        for kind in ('csv','xlsx'):atomic_bytes(folder/f'transactions.{kind}',export_bytes(kind,prepared,document))
        source=Path(data)/'documents'/document['id']/'searchable.pdf'
        target=folder/'searchable-statement.pdf'
        if document['report'].get('searchable') and source.is_file() and not target.exists():
            try:os.link(source,target)
            except OSError:
                import shutil
                shutil.copyfile(source,target)
    return folder
